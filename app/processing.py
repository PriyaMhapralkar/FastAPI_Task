"""Text extraction, cleaning and chunking. Pure Python, no database, no LLM."""

import logging
import re
import unicodedata
from pathlib import Path

from docx import Document as DocxDocument
from langchain_text_splitters import RecursiveCharacterTextSplitter
from pypdf import PdfReader

logger = logging.getLogger(__name__)

CHUNK_SIZE = 500
CHUNK_OVERLAP = 50


class ProcessingError(Exception):
    """Raised when a document cannot be turned into chunks."""



def extract_text(path: Path, extension: str) -> str:
    """Read a file from disk and return its raw text."""

    try:
        if extension == ".pdf":
            reader = PdfReader(str(path))
            return "\n".join(
                page.extract_text() or ""
                for page in reader.pages
            )

        if extension == ".docx":
            doc = DocxDocument(str(path))

            parts = [p.text for p in doc.paragraphs]

            for table in doc.tables:
                for row in table.rows:
                    parts.append(
                        " ".join(cell.text for cell in row.cells)
                    )

            return "\n".join(parts)

        if extension in {".txt", ".md"}:
            return path.read_text(
                encoding="utf-8-sig",
                errors="replace"
            )

    except Exception as exc:
        logger.exception("Could not read %s", path)

        raise ProcessingError(
            f"Could not read the {extension} file. "
            "It may be corrupt or protected."
        ) from exc

    raise ProcessingError(
        f"Text extraction is not supported for '{extension}' files."
    )



_CONTROL_CHARS = re.compile(
    r"[\x00-\x08\x0b-\x1f\x7f\u200b\u200c\u200d\ufeff]"
)


def clean_text(text: str) -> str:
    """
    Remove unwanted characters and repeated spaces,
    while keeping paragraph and line breaks.
    """

    text = unicodedata.normalize("NFKC", text)

    text = text.replace("\r\n", "\n").replace("\r", "\n")

    text = _CONTROL_CHARS.sub("", text)

    text = re.sub(r"[ \t\f\v]+", " ", text)

    text = re.sub(r" ?\n ?", "\n", text)

    text = re.sub(r"\n{3,}", "\n\n", text)

    return text.strip()



_splitter = RecursiveCharacterTextSplitter(
    chunk_size=CHUNK_SIZE,
    chunk_overlap=CHUNK_OVERLAP,
    separators=["\n\n", "\n", " ", ""],
)


def chunk_text(text: str) -> list[str]:
    """Split cleaned text into chunks of at most CHUNK_SIZE characters."""

    chunks: list[str] = []

    for piece in _splitter.split_text(text):

        flat = re.sub(r"\s+", " ", piece).strip()

        if not flat:
            continue

        for start in range(0, len(flat), CHUNK_SIZE):
            chunk = flat[start:start + CHUNK_SIZE].strip()

            if chunk:
                chunks.append(chunk)

    return chunks